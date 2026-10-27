#!/usr/bin/env python3
"""Build the Office fixtures the office-viewer e2e spec opens.

Each file carries the product feature the viewer must keep visible:
  review.docx  two comments by two reviewers and a tracked insertion
  deck.pptx    two slides, speaker notes on both, a comment on slide 2
  budget.xlsx  a cell comment (note) on B2
Run from this directory: python3 make-office-fixtures.py
"""
import io, re, zipfile
from datetime import datetime

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl
from openpyxl.comments import Comment
from pptx import Presentation
from pptx.util import Inches, Pt


def make_docx(path):
    d = docx.Document()
    d.add_heading('Quarterly review', level=1)
    p = d.add_paragraph('Revenue grew 18 percent on the strength of the retrieval product.')
    run = p.runs[0]
    c = d.add_comment(run, text='Cite the source for 18 percent.', author='Dana Reviewer', initials='DR')
    d.add_paragraph('Churn held flat at 2 percent.')
    p2 = d.add_paragraph('Next quarter we expand to three regions.')
    d.add_comment(p2.runs[0], text='Which three regions?', author='Lee Editor', initials='LE')
    # A tracked insertion (w:ins) the viewer must show as a change.
    p3 = d.add_paragraph('Hiring stays ')
    ins = OxmlElement('w:ins')
    ins.set(qn('w:id'), '901'); ins.set(qn('w:author'), 'Lee Editor'); ins.set(qn('w:date'), '2026-10-27T12:00:00Z')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'deliberately '; t.set(qn('xml:space'), 'preserve')
    r.append(t); ins.append(r); p3._p.append(ins)
    tail = OxmlElement('w:r'); tt = OxmlElement('w:t'); tt.text = 'slow.'; tail.append(tt); p3._p.append(tail)
    d.save(path)


LEGACY_AUTHORS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:cmAuthor id="0" name="Sam Presenter" initials="SP" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>')
LEGACY_COMMENTS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:cm authorId="0" dt="2026-10-27T12:00:00.000" idx="1"><p:pos x="10" y="10"/>'
    '<p:text>Swap this chart for the October numbers.</p:text></p:cm></p:cmLst>')
REL_AUTHORS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors'
REL_COMMENTS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments'


def make_pptx(path):
    prs = Presentation()
    for i, (title, body, notes) in enumerate([
            ('Retrieval at scale', 'One index for video, images and text', 'Open with the customer story.'),
            ('October results', 'Latency down 40 percent', 'Pause here for questions.')]):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO(); prs.save(buf)
    # python-pptx writes no comments: add a legacy comment on slide 2 by hand.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED)
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename == '[Content_Types].xml':
            data = data.replace(b'</Types>',
                b'<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/>'
                b'<Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/></Types>')
        elif item.filename == 'ppt/_rels/presentation.xml.rels':
            data = data.replace(b'</Relationships>',
                ('<Relationship Id="rIdCmA" Type="%s" Target="commentAuthors.xml"/></Relationships>' % REL_AUTHORS).encode())
        elif item.filename == 'ppt/slides/_rels/slide2.xml.rels':
            data = data.replace(b'</Relationships>',
                ('<Relationship Id="rIdCm1" Type="%s" Target="../comments/comment1.xml"/></Relationships>' % REL_COMMENTS).encode())
        out.writestr(item, data)
    out.writestr('ppt/commentAuthors.xml', LEGACY_AUTHORS)
    out.writestr('ppt/comments/comment1.xml', LEGACY_COMMENTS)
    out.close()


def make_xlsx(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Budget'
    ws['A1'], ws['B1'] = 'Item', 'Cost'
    ws['A2'], ws['B2'] = 'GPUs', 42000
    ws['A3'], ws['B3'] = 'Storage', 3100
    ws['B2'].comment = Comment('Quote from the October vendor call.', 'Pat Finance')
    wb.save(path)


if __name__ == '__main__':
    make_docx('review.docx'); make_pptx('deck.pptx'); make_xlsx('budget.xlsx')
    print('wrote review.docx deck.pptx budget.xlsx')
